"""查询与导出（Excel / CSV）。"""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import String, cast, or_, select

from .config import scoring
from .db import Store, User, session

COLUMNS: list[tuple[str, str, int]] = [
    ("domain", "网址", 28),
    ("name", "店铺名", 22),
    ("grade", "等级", 6),
    ("total_score", "总分", 7),
    ("main_products", "主营产品", 36),
    ("main_categories", "主营品类", 24),
    ("promotion_channels", "推广方式", 30),
    ("chinese_type", "华人类型", 12),
    ("chinese_confidence", "华人置信度", 10),
    ("cn_rule_score", "华人规则分", 10),
    ("platform", "建站平台", 11),
    ("country", "国家", 6),
    ("location", "所在地", 20),
    ("product_count", "商品数", 8),
    ("price_range", "价格区间", 14),
    ("emails", "邮箱", 30),
    ("phones", "电话", 20),
    ("socials", "社媒", 40),
    ("outreach_angle", "建联切入点", 40),
    ("opening_line", "英文开场白", 40),
    ("reasoning", "打分理由", 50),
    ("scores", "分项得分", 40),
    ("cn_evidence", "华人证据", 50),
    ("outreach_status", "跟进状态", 10),
    ("owner", "负责人", 10),
    ("notes", "备注", 30),
    ("source", "发现渠道", 12),
    ("discovered_at", "发现时间", 18),
    ("scored_at", "打分时间", 18),
]


def _dim_label(key: str) -> str:
    return scoring().get("dimensions", {}).get(key, {}).get("label", key)


def user_names() -> dict[int, str]:
    with session() as s:
        return {u.id: u.display_name or u.username for u in s.scalars(select(User))}


def store_row(st: Store, users: dict[int, str] | None = None) -> dict[str, Any]:
    r = st.llm_result or {}
    out = r.get("outreach") or {}
    cn = r.get("chinese") or {}
    price = ""
    if st.price_min is not None:
        price = f"${st.price_min:.0f}-${st.price_max:.0f}" if st.price_max else f"${st.price_min:.0f}"
    return {
        "id": st.id,
        "domain": st.domain,
        "url": st.url or f"https://{st.domain}",
        "name": st.name,
        "status": st.status,
        "grade": st.grade,
        "total_score": st.total_score,
        "main_products": st.main_products,
        "main_categories": "、".join(st.main_categories or []),
        "promotion_channels": "、".join(st.promotion_channels or []),
        "chinese_type": st.chinese_type,
        "chinese_confidence": st.chinese_confidence,
        "cn_rule_score": st.cn_rule_score,
        "platform": st.platform,
        "country": st.country,
        "location": ", ".join(x for x in (st.city, st.province) if x),
        "product_count": st.product_count,
        "price_range": price,
        "emails": ", ".join(st.emails or []),
        "phones": ", ".join(st.phones or []),
        "socials": " ".join((st.socials or {}).values()),
        "outreach_angle": out.get("angle"),
        "opening_line": out.get("opening_line"),
        "reasoning": r.get("reasoning"),
        "scores": ", ".join(f"{_dim_label(k)}:{v}" for k, v in (st.scores or {}).items()),
        "cn_evidence": "; ".join(cn.get("evidence") or [s.get("detail", "") for s in st.cn_signals or []]),
        "outreach_status": st.outreach_status,
        "owner_id": st.owner_id,
        "owner": (users or {}).get(st.owner_id) if st.owner_id else None,
        "starred": bool(st.starred),
        "notes": st.notes,
        "source": st.source,
        "discovered_at": st.discovered_at.strftime("%Y-%m-%d %H:%M") if st.discovered_at else "",
        "scored_at": st.scored_at.strftime("%Y-%m-%d %H:%M") if st.scored_at else "",
    }


def build_query(*, q: str | None = None, status: str | None = None, grade: str | None = None,
                chinese_type: str | None = None, platform: str | None = None, min_score: float | None = None,
                outreach_status: str | None = None, sort: str = "score", owner_id: int | None = None,
                unowned: bool = False, starred: bool = False, category: str | None = None):
    stmt = select(Store)
    if owner_id:
        stmt = stmt.where(Store.owner_id == owner_id)
    if unowned:
        stmt = stmt.where(Store.owner_id.is_(None))
    if starred:
        stmt = stmt.where(Store.starred.is_(True))
    if category:
        # JSON 列在 SQLite/Postgres 中都可以按文本模糊匹配
        stmt = stmt.where(cast(Store.main_categories, String).ilike(f"%{category}%"))
    # 默认只看已打分；status="all" 表示不过滤
    if not status:
        stmt = stmt.where(Store.status == "scored")
    elif status != "all":
        stmt = stmt.where(Store.status.in_(status.split(",")))
    if grade:
        stmt = stmt.where(Store.grade.in_(grade.split(",")))
    if chinese_type:
        stmt = stmt.where(Store.chinese_type.in_(chinese_type.split(",")))
    if platform:
        stmt = stmt.where(Store.platform == platform)
    if min_score:
        stmt = stmt.where(Store.total_score >= min_score)
    if outreach_status:
        stmt = stmt.where(Store.outreach_status == outreach_status)
    if q:
        like = f"%{q}%"
        stmt = stmt.where(or_(Store.domain.ilike(like), Store.name.ilike(like), Store.main_products.ilike(like)))
    order = {
        "score": [Store.total_score.desc().nulls_last()],
        "cn": [Store.cn_rule_score.desc().nulls_last()],
        "new": [Store.discovered_at.desc()],
        "scored": [Store.scored_at.desc().nulls_last()],
        "updated": [Store.updated_at.desc().nulls_last()],
    }.get(sort, [Store.total_score.desc().nulls_last()])
    return stmt.order_by(*order, Store.id.desc())


def query_stores(limit: int = 100000, **filters) -> list[dict]:
    users = user_names()
    with session() as s:
        return [store_row(st, users) for st in s.scalars(build_query(**filters).limit(limit))]


def export_rows(rows: list[dict], path: Path) -> None:
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with path.open("w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow([c[1] for c in COLUMNS])
            for r in rows:
                w.writerow([r.get(c[0]) for c in COLUMNS])
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "独立站"
    ws.append([c[1] for c in COLUMNS])
    for r in rows:
        ws.append([r.get(c[0]) for c in COLUMNS])
    head_fill = PatternFill("solid", fgColor="1F4E78")
    for i, (_, _, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = head_fill
        ws.column_dimensions[cell.column_letter].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].hyperlink = f"https://{row[0].value}"
        row[0].font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
